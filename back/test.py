from openpyxl import load_workbook, Workbook
wb = load_workbook("brands/Apple/iPhone.xlsx")
ws = wb["Дисплей"]
raise
row = 2
ws.cell(row=row, column=4).value = 123123
wb.save("brands/Apple/iPhone.xlsx")